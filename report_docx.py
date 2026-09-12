"""Phase 5.3 — editable DOCX compliance report generation.

Same data source as report_pdf.py (build_report_data()) so both formats
always agree on content — only the rendering differs. This produces a
genuinely editable Word document (real headings/paragraphs/table an
officer can select and type into), not a PDF wrapped in a docx container.
"""

import io
from typing import TYPE_CHECKING

from docx import Document
from docx.shared import Pt, RGBColor

from reporting import build_report_data

if TYPE_CHECKING:
    from models import Scan


STATUS_COLORS = {
    "COMPLIANT": RGBColor(0x1A, 0x7F, 0x37),
    "WARNING": RGBColor(0x9A, 0x67, 0x00),
    "VIOLATION": RGBColor(0xCF, 0x22, 0x2E),
}
DEFAULT_STATUS_COLOR = RGBColor(0x57, 0x60, 0x6A)


def _add_status_paragraph(document: Document, status: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(status)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = STATUS_COLORS.get(status, DEFAULT_STATUS_COLOR)


def _add_field_table(document: Document, declared_fields: list[dict]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"

    for field in declared_fields:
        row = table.add_row().cells
        row[0].text = field["label"]
        row[1].text = field["value"]


def _add_issue_list(document: Document, issues: list[str]) -> None:
    for issue in issues:
        document.add_paragraph(issue, style="List Bullet")


def generate_docx_report(scan: "Scan") -> bytes:
    """Build the compliance report for `scan` and return raw .docx bytes."""
    data = build_report_data(scan)
    document = Document()

    document.add_heading("SetuCheck Legal Metrology Compliance Report", level=1)

    subtitle = document.add_paragraph()
    subtitle_run = subtitle.add_run(
        f"Product: {data['product_name']}  \u00b7  Category: {data['category']}"
    )
    subtitle_run.font.color.rgb = DEFAULT_STATUS_COLOR

    _add_status_paragraph(document, data["status"])

    document.add_heading("Declared Fields", level=2)
    _add_field_table(document, data["declared_fields"])

    document.add_heading("Label Quality Checks", level=2)
    quality_table = document.add_table(rows=0, cols=2)
    quality_table.style = "Light Grid Accent 1"
    for label, value in (
        ("Readability", data["readability"]),
        ("Placement", data["placement"]),
    ):
        row = quality_table.add_row().cells
        row[0].text = label
        row[1].text = value

    document.add_heading("Issues", level=2)
    _add_issue_list(document, data["issues"])

    meta = document.add_paragraph()
    meta_run = meta.add_run(
        f"Scan ID: {data['scan_id']}  \u00b7  "
        f"Input type: {data['input_type']}  \u00b7  "
        f"Scanned: {data['created_at']}"
    )
    meta_run.font.size = Pt(8)
    meta_run.font.color.rgb = DEFAULT_STATUS_COLOR

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()
